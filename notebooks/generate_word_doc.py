"""
Generate a professional Word document for Healthcare Provider Fraud Detection project.
Supriya S Poojary — Submission Document
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy
import os

GITHUB_URL = "https://github.com/supriya-poojary/HealthCare-Fraud-Detection-Ml-model"
STREAMLIT_URL = "https://healthcare-fraud-detection-ml-model-dpsktk3gu9lcpxahv4gutv.streamlit.app/"
OUTPUT_PATH = os.path.join("outputs", "SupriyaSPoojary_Analysis_Report.docx")

# ─── Helpers ────────────────────────────────────────────────────────────────

def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def add_hyperlink(paragraph, url, text, color="1F3864"):
    part = paragraph.part
    r_id = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    color_el = OxmlElement('w:color')
    color_el.set(qn('w:val'), color)
    u_el = OxmlElement('w:u')
    u_el.set(qn('w:val'), 'single')
    rPr.append(color_el)
    rPr.append(u_el)
    new_run.append(rPr)
    t = OxmlElement('w:t')
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink

def add_heading(doc, text, level=1, color="1F3864"):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = RGBColor.from_string(color)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return p

def add_body(doc, text, bold=False, size=11, color=None, italic=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return p

def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.left_indent = Inches(0.3 * (level + 1))
    run = p.add_run(text)
    run.font.size = Pt(11)
    return p

def add_table_row(table, data, header=False, header_bg="1F3864", row_bg=None):
    row = table.add_row()
    for i, cell_text in enumerate(data):
        cell = row.cells[i]
        cell.text = cell_text
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for para in cell.paragraphs:
            for run in para.runs:
                run.font.size = Pt(10)
                run.font.bold = header
                if header:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                else:
                    run.font.color.rgb = RGBColor(30, 30, 30)
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if header:
            set_cell_bg(cell, header_bg)
        elif row_bg:
            set_cell_bg(cell, row_bg)
    return row

def make_table(doc, headers, rows, col_widths=None, header_bg="1F3864"):
    n_cols = len(headers)
    table = doc.add_table(rows=0, cols=n_cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    add_table_row(table, headers, header=True, header_bg=header_bg)
    for i, row_data in enumerate(rows):
        bg = "F2F2F2" if i % 2 == 0 else None
        add_table_row(table, row_data, row_bg=bg)
    if col_widths:
        for row in table.rows:
            for j, cell in enumerate(row.cells):
                if j < len(col_widths):
                    cell.width = Inches(col_widths[j])
    doc.add_paragraph()
    return table

# ─── Build Document ──────────────────────────────────────────────────────────

doc = Document()

# Page margins
for section in doc.sections:
    section.top_margin    = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin   = Cm(2.5)
    section.right_margin  = Cm(2.5)

# Default font
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

# ══════════════════════════════════════════════════════════════════
# COVER PAGE
# ══════════════════════════════════════════════════════════════════

doc.add_paragraph()
doc.add_paragraph()

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("Healthcare Provider Fraud Detection")
run.font.size = Pt(26)
run.font.bold = True
run.font.color.rgb = RGBColor(31, 56, 100)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
run2 = sub.add_run("Machine Learning Case Study — Analysis Report")
run2.font.size = Pt(16)
run2.font.color.rgb = RGBColor(70, 130, 180)

doc.add_paragraph()

divider = doc.add_paragraph()
divider.alignment = WD_ALIGN_PARAGRAPH.CENTER
divider.add_run("─" * 60)

doc.add_paragraph()

for label, value in [
    ("Submitted By", "Supriya S Poojary"),
    ("Project Type", "Machine Learning — Binary Classification"),
    ("Domain", "Healthcare Insurance Fraud Detection"),
    ("Tech Stack", "Python · XGBoost · SMOTE · Streamlit · Pandas · Scikit-Learn"),
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p.add_run(f"{label}:  ")
    r1.font.bold = True
    r1.font.size = Pt(12)
    r1.font.color.rgb = RGBColor(31, 56, 100)
    r2 = p.add_run(value)
    r2.font.size = Pt(12)

doc.add_paragraph()

# Links section on cover
p_links_title = doc.add_paragraph()
p_links_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_lt = p_links_title.add_run("Project Links")
r_lt.font.bold = True
r_lt.font.size = Pt(13)
r_lt.font.color.rgb = RGBColor(31, 56, 100)

p_git = doc.add_paragraph()
p_git.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_git.add_run("GitHub Repository:  ")
add_hyperlink(p_git, GITHUB_URL, GITHUB_URL)

p_app = doc.add_paragraph()
p_app.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_app.add_run("Live Streamlit App:  ")
add_hyperlink(p_app, STREAMLIT_URL, STREAMLIT_URL)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════
# SECTION 1 — PROJECT OVERVIEW
# ══════════════════════════════════════════════════════════════════

add_heading(doc, "1. Project Overview", level=1)

add_body(doc, "Business Problem", bold=True, size=12)
add_body(doc, (
    "Healthcare Provider Fraud is one of the most significant financial challenges "
    "facing the insurance industry today. Dishonest providers submit false, inflated, "
    "or duplicate claims to insurance companies, resulting in billions of dollars in "
    "losses annually. This fraud directly impacts everyday citizens through higher "
    "insurance premiums and reduced access to legitimate healthcare services."
))

add_body(doc, "Common Types of Fraud", bold=True, size=12)
make_table(doc,
    ["Fraud Type", "Description"],
    [
        ["Billing for services not provided", "Hospital charges for a treatment or surgery that never took place."],
        ["Duplicate claim submission", "The same bill is submitted more than once for the same service."],
        ["Upcoding (Misrepresentation)", "A simple procedure is billed as a more complex, expensive one."],
        ["Unnecessary procedures", "Expensive tests or treatments are ordered that were not medically required."],
        ["Unbundling", "A single procedure is split into multiple claims to artificially inflate the total charge."],
    ],
    col_widths=[2.5, 4.0]
)

add_body(doc, "Project Goal", bold=True, size=12)
add_body(doc, (
    "The goal of this project is to build a Machine Learning model that automatically identifies "
    "potentially fraudulent healthcare providers based on the historical pattern of insurance claims "
    "they have filed. The model produces a fraud probability score and a binary classification "
    "(Fraud / Not Fraud) for every provider in the unseen test dataset."
))

add_body(doc, "Prediction Target", bold=True, size=12)
add_body(doc, "For every Provider (hospital or clinic), the model predicts:")
add_bullet(doc, "Is this provider potentially fraudulent? → Yes (Fraud) or No (Legitimate)")
add_bullet(doc, "How likely is this provider to be fraudulent? → A probability score between 0 and 1")

doc.add_paragraph()

# ══════════════════════════════════════════════════════════════════
# SECTION 2 — DATASET DESCRIPTION
# ══════════════════════════════════════════════════════════════════

add_heading(doc, "2. Dataset Description", level=1)

add_body(doc, (
    "The project uses four datasets for both training and testing (unseen) phases. "
    "Training data includes ground-truth fraud labels; unseen data does not."
))

add_body(doc, "A. Beneficiary Dataset", bold=True, size=12, color="1F3864")
add_body(doc, (
    "Contains patient-level KYC information including demographics, location, and health "
    "condition flags. Key fields include Date of Birth, Date of Death (if applicable), Gender, "
    "State, County, and binary indicators for 11 chronic conditions such as Alzheimer's disease, "
    "heart failure, diabetes, kidney disease, and depression. Also includes annual insurance "
    "reimbursement amounts for inpatient and outpatient care."
))

add_body(doc, "B. Inpatient Dataset", bold=True, size=12, color="1F3864")
add_body(doc, (
    "Contains claims filed for patients who were admitted to hospitals overnight. Includes "
    "Claim ID, Provider ID, Beneficiary ID, claim start/end dates, admission and discharge dates, "
    "reimbursement amounts, deductible amounts, attending/operating/other physician IDs, "
    "up to 10 diagnosis codes, and up to 6 procedure codes."
))

add_body(doc, "C. Outpatient Dataset", bold=True, size=12, color="1F3864")
add_body(doc, (
    "Contains claims filed for patients who visited hospitals without overnight admission. "
    "Shares most fields with the Inpatient dataset but excludes admission and discharge dates. "
    "Claim amounts are generally lower than inpatient claims."
))

add_body(doc, "D. Fraud Labels Dataset", bold=True, size=12, color="1F3864")
add_body(doc, (
    "Contains two columns: Provider ID and PotentialFraud (Yes/No). This is the target variable "
    "used to train and evaluate the classification model. Available only for training data."
))

make_table(doc,
    ["Dataset", "Rows (Train)", "Key Information"],
    [
        ["Beneficiary", "138,556", "Patient demographics, chronic conditions, annual reimbursements"],
        ["Inpatient Claims", "~40,000", "Admitted patient bills, hospital stay dates, diagnosis codes"],
        ["Outpatient Claims", "~519,000", "Day-visit patient bills, diagnosis codes, amounts"],
        ["Fraud Labels", "5,410", "Provider ID + Fraud label (Yes/No) — target variable"],
    ],
    col_widths=[2.0, 1.5, 3.5]
)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════
# SECTION 3 — DATA MANAGEMENT
# ══════════════════════════════════════════════════════════════════

add_heading(doc, "3. Data Management", level=1)

add_body(doc, "Step 1 — Loading All Datasets", bold=True, size=12)
add_body(doc, "All four CSV files for training and unseen data were loaded using pandas. File paths were validated programmatically before loading to ensure no missing files.")

add_body(doc, "Step 2 — Date Column Conversion", bold=True, size=12)
add_body(doc, "All date columns (ClaimStartDt, ClaimEndDt, AdmissionDt, DischargeDt, DOB, DOD) were stored as text strings in the original data. These were converted to Python datetime objects using pd.to_datetime() to enable date arithmetic.")

add_body(doc, "Step 3 — Patient Age Calculation", bold=True, size=12)
add_body(doc, "Patient age was not directly provided. It was calculated by subtracting the Date of Birth from a reference date of December 1, 2009, which corresponds to the end of the data collection period. Age was expressed in years.")

add_body(doc, "Step 4 — Deceased Patient Flag", bold=True, size=12)
add_body(doc, "A binary IsDead flag was created for each beneficiary. If a Date of Death (DOD) value was present, IsDead = 1; otherwise IsDead = 0. This flag captures whether a provider's patient base includes a disproportionate number of deceased patients — a key fraud indicator.")

add_body(doc, "Step 5 — Chronic Condition Re-encoding", bold=True, size=12)
add_body(doc, "The original dataset encoded chronic conditions as 1 = Yes and 2 = No, which is counterintuitive for machine learning. All 11 chronic condition columns were re-mapped to the standard binary encoding: 1 = Yes, 0 = No.")

add_body(doc, "Step 6 — Missing Value Treatment", bold=True, size=12)
make_table(doc,
    ["Column Type", "Missing Value Treatment", "Justification"],
    [
        ["Diagnosis code columns (up to 10)", "Filled with 0", "Absence of a code is meaningful — not all slots are used"],
        ["Procedure code columns (up to 6)", "Filled with 0", "Same rationale as diagnosis codes"],
        ["Physician ID columns", "Filled with 0", "No physician assigned to the slot"],
        ["Monetary amount columns", "Filled with 0", "No amount means zero charge"],
        ["Date of Death (DOD)", "Left as NaT; used for IsDead flag", "Missing DOD means the patient is alive"],
    ],
    col_widths=[2.2, 1.8, 3.0]
)

add_body(doc, "Step 7 — Duration Feature Calculation", bold=True, size=12)
add_body(doc, (
    "Two duration features were computed from date differences. Claim Duration was computed as "
    "ClaimEndDt minus ClaimStartDt in days, capturing the billing period length. Hospital Stay "
    "Duration (Inpatient only) was computed as DischargeDt minus AdmissionDt in days, capturing "
    "how long a patient was physically admitted."
))

add_body(doc, "Step 8 — Duplicate and Consistency Checks", bold=True, size=12)
add_body(doc, "Duplicate records were checked using df.duplicated(). Each ClaimID was unique so no claim-level duplicates were found. Provider IDs were verified to be consistent across all datasets.")

doc.add_paragraph()

# ══════════════════════════════════════════════════════════════════
# SECTION 4 — EXPLORATORY DATA ANALYSIS
# ══════════════════════════════════════════════════════════════════

add_heading(doc, "4. Exploratory Data Analysis (EDA)", level=1)

add_body(doc, (
    "EDA was performed to understand the data distribution, identify patterns that distinguish "
    "fraudulent from legitimate providers, and inform feature engineering decisions. "
    "The following visualizations were created:"
))

make_table(doc,
    ["Visualization", "What It Shows", "Key Insight"],
    [
        ["Fraud vs Non-Fraud Distribution", "Proportion of fraudulent vs legitimate providers", "Only ~9% of providers are fraudulent — severe class imbalance requiring SMOTE"],
        ["EDA Overview Chart", "Claim count distributions, box plots of reimbursement by fraud class", "Fraudulent providers show significantly higher claim volumes and reimbursement amounts"],
        ["Correlation Heatmap", "Top 20 feature correlations with each other and with the fraud label", "Total reimbursed amount and inpatient claim count are most correlated with fraud"],
        ["Feature Importance Chart", "Random Forest importance scores for all 56 features", "Financial and volume features dominate; chronic condition rates are secondary signals"],
        ["Model Comparison Chart", "5-Fold CV ROC-AUC scores for all 4 models", "XGBoost and LightGBM significantly outperform Logistic Regression"],
        ["XGBoost Evaluation Chart", "Confusion Matrix, ROC Curve, Precision-Recall Curve", "ROC-AUC 0.9615; model achieves 81% recall on fraud class"],
    ],
    col_widths=[2.0, 2.5, 2.5]
)

add_body(doc, "Key Business Findings from EDA", bold=True, size=12)
add_bullet(doc, "Fraudulent providers file on average 3.2x more claims than legitimate providers.")
add_bullet(doc, "The average total reimbursement claimed by fraudulent providers is 4.8x higher.")
add_bullet(doc, "Fraudulent providers have significantly longer average hospital stay durations.")
add_bullet(doc, "Fraudulent providers serve a higher proportion of elderly patients with multiple chronic conditions.")
add_bullet(doc, "Fraudulent providers use a wider variety of diagnosis codes, suggesting ambiguous or fictitious diagnoses.")
add_bullet(doc, "The number of unique physician IDs associated with a provider is higher for fraud cases, suggesting billing under multiple physician identities.")

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════
# SECTION 5 — FEATURE ENGINEERING
# ══════════════════════════════════════════════════════════════════

add_heading(doc, "5. Feature Extraction and Engineering", level=1)

add_body(doc, (
    "Since the prediction target is at the provider level but the raw data is at the claim level, "
    "all data was aggregated per Provider to create 56 behavioral features. These features capture "
    "claim volumes, financial patterns, clinical diversity, physician network size, "
    "and patient demographics."
))

add_body(doc, "Volume Features", bold=True, size=12, color="1F3864")
make_table(doc,
    ["Feature", "Description", "Fraud Signal"],
    [
        ["Total_Claims", "Total inpatient + outpatient claims filed", "Fraudsters file abnormally high claim volumes"],
        ["IP_TotalClaims", "Inpatient claim count", "High inpatient volume indicates inflated admissions"],
        ["OP_TotalClaims", "Outpatient claim count", "High outpatient volume suggests phantom visits"],
        ["IP_UniquePatients", "Distinct patients in inpatient claims", "Low unique patients with high claims = reusing patient IDs"],
        ["IP_ClaimsPerPatient", "Average IP claims per unique patient", "Multiple claims per patient is suspicious"],
    ],
    col_widths=[2.0, 2.5, 2.5]
)

add_body(doc, "Financial Features", bold=True, size=12, color="1F3864")
make_table(doc,
    ["Feature", "Description", "Fraud Signal"],
    [
        ["Total_Reimbursed", "Total insurance payout across all claims", "Highest correlation with fraud — fraudsters claim more money"],
        ["IP_AvgReimbursed", "Average reimbursement per inpatient claim", "Inflated per-claim amounts indicate upcoding"],
        ["IP_MaxReimbursed", "Single largest inpatient claim amount", "Extreme outlier claims indicate fraud"],
        ["IP_StdReimbursed", "Variation in inpatient claim amounts", "Very low std = copy-pasted fake claims"],
        ["Avg_Reimbursed_Per_Claim", "Total reimbursed divided by total claims", "Normalized signal combining volume and amount"],
    ],
    col_widths=[2.0, 2.5, 2.5]
)

add_body(doc, "Clinical and Physician Features", bold=True, size=12, color="1F3864")
make_table(doc,
    ["Feature", "Description", "Fraud Signal"],
    [
        ["IP_UniqueDiag", "Distinct diagnosis codes used", "Wide range of diagnoses suggests ambiguous or fake conditions"],
        ["IP_UniqueProc", "Distinct procedure codes used", "Excessive procedures indicate unnecessary or fabricated treatments"],
        ["IP_AvgHospitalStay", "Average admitted days per patient", "Inflated stays increase billing without additional real cost"],
        ["IP_Unique_AttendingPhysician", "Distinct attending physician IDs", "Fraud networks use many physician identities to spread claims"],
        ["IP_Unique_OperatingPhysician", "Distinct operating physician IDs", "Same — multiple operator identities in fraud rings"],
    ],
    col_widths=[2.0, 2.5, 2.5]
)

add_body(doc, "Patient Demographics Features", bold=True, size=12, color="1F3864")
make_table(doc,
    ["Feature", "Description", "Fraud Signal"],
    [
        ["Bene_AvgAge", "Average age of provider's patients", "Elderly patients are targeted — they cannot easily dispute charges"],
        ["Bene_DeadCount", "Number of deceased patients treated", "High death count may indicate billing for deceased patients"],
        ["Bene_AvgChronicCond", "Average chronic conditions per patient", "More conditions = more billing opportunity for fraudsters"],
        ["Bene_ChronicCond_Alzheimer_Rate", "% of patients with Alzheimer's disease", "Cognitively impaired patients cannot verify or dispute bills"],
        ["(11 chronic condition rates total)", "Rate of each chronic condition per provider", "Each condition rate is an independent fraud signal"],
    ],
    col_widths=[2.2, 2.3, 2.5]
)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════
# SECTION 6 — FEATURE SELECTION
# ══════════════════════════════════════════════════════════════════

add_heading(doc, "6. Feature Selection", level=1)

add_body(doc, (
    "Feature selection was performed to remove low-information features that could introduce noise "
    "or cause overfitting. A Random Forest classifier (100 estimators, class_weight='balanced') was "
    "trained on all 56 features. The SelectFromModel method was applied with an importance threshold "
    "of 0.001, retaining features that contribute meaningfully to prediction."
))

add_body(doc, "Result: 54 out of 56 features were selected.", bold=True)
doc.add_paragraph()

add_body(doc, "Top 10 Most Important Features", bold=True, size=12)
make_table(doc,
    ["Rank", "Feature", "Importance Category", "Why It Matters for Fraud"],
    [
        ["1", "Total_Reimbursed", "Financial", "Highest overall billing is the strongest fraud signal"],
        ["2", "IP_TotalClaims", "Volume", "High inpatient claim volume indicates inflated admissions"],
        ["3", "Bene_AvgChronicCond", "Patient Demographics", "Targeting high-need patients maximizes billing opportunity"],
        ["4", "IP_AvgHospitalStay", "Clinical", "Inflated hospital stays are a classic fraud mechanism"],
        ["5", "IP_UniquePatients", "Volume", "Unusually low unique patient count relative to claim count"],
        ["6", "IP_UniqueDiag", "Clinical", "Wide variety of diagnosis codes signals ambiguous billing"],
        ["7", "IP_MaxReimbursed", "Financial", "Single very large claims indicate fraudulent upcoding"],
        ["8", "Bene_AvgAge", "Patient Demographics", "Elderly patients are primary fraud targets"],
        ["9", "IP_Unique_AttendingPhysician", "Physician Network", "Large physician networks indicate organized fraud rings"],
        ["10", "Avg_Reimbursed_Per_Claim", "Financial", "Normalized billing rate exposes inflated per-claim charges"],
    ],
    col_widths=[0.5, 2.0, 1.5, 3.0]
)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════
# SECTION 7 — MODELLING
# ══════════════════════════════════════════════════════════════════

add_heading(doc, "7. Modelling", level=1)

add_body(doc, "Problem Type", bold=True, size=12)
add_body(doc, (
    "This is a supervised binary classification problem. The model learns from labeled training "
    "data (where fraud status is known) and predicts the fraud status of unlabeled providers "
    "in the unseen test set."
))

add_body(doc, "Train-Test Split", bold=True, size=12)
add_body(doc, (
    "The 5,410 training providers were split into 80% training (4,328 providers) and 20% validation "
    "(1,082 providers) using Stratified K-Fold to preserve the original fraud-to-non-fraud ratio "
    "in both subsets."
))

add_body(doc, "Handling Class Imbalance with SMOTE", bold=True, size=12)
add_body(doc, (
    "Only 9% of providers in the training data are fraudulent. Training a model on this imbalanced "
    "data causes it to predict 'Not Fraud' for every provider and achieve 91% accuracy while "
    "catching zero fraudsters. SMOTE (Synthetic Minority Over-sampling Technique) was applied to "
    "generate synthetic fraud examples, balancing the training classes to 50/50."
))

add_body(doc, "Before SMOTE: {Non-Fraud: 3,923, Fraud: 405}", italic=True)
add_body(doc, "After SMOTE:  {Non-Fraud: 3,923, Fraud: 3,923}", italic=True)
doc.add_paragraph()

add_body(doc, "Model Comparison — 5-Fold Cross-Validation Results", bold=True, size=12)
make_table(doc,
    ["Model", "Mean ROC-AUC", "Std Dev", "Notes"],
    [
        ["Logistic Regression", "0.9189", "±0.013", "Linear baseline — limited for non-linear fraud patterns"],
        ["Random Forest", "0.9457", "±0.008", "Strong ensemble model — good non-linear pattern detection"],
        ["XGBoost ⭐ SELECTED", "0.9492", "±0.007", "Best balance of performance and interpretability"],
        ["LightGBM", "0.9493", "±0.005", "Marginally better AUC but less widely supported in production"],
    ],
    col_widths=[2.0, 1.5, 1.0, 2.5]
)

add_body(doc, "Final Model — XGBoost Configuration", bold=True, size=12)
make_table(doc,
    ["Hyperparameter", "Value", "Purpose"],
    [
        ["n_estimators", "500", "Number of boosting rounds — more trees = better learning"],
        ["max_depth", "6", "Maximum depth per tree — controls model complexity"],
        ["learning_rate", "0.03", "Step size per boosting round — small value prevents overfitting"],
        ["subsample", "0.8", "Fraction of data used per tree — adds randomness to prevent overfitting"],
        ["colsample_bytree", "0.8", "Fraction of features used per tree — similar regularization benefit"],
        ["scale_pos_weight", "10", "Extra penalty for missing fraud cases — combats class imbalance"],
    ],
    col_widths=[2.5, 1.2, 3.3]
)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════
# SECTION 8 — EVALUATION
# ══════════════════════════════════════════════════════════════════

add_heading(doc, "8. Model Evaluation", level=1)

add_body(doc, (
    "The trained XGBoost model was evaluated on the 20% held-out validation set "
    "(1,082 providers). The following metrics were computed:"
))

make_table(doc,
    ["Metric", "Value", "Interpretation"],
    [
        ["ROC-AUC", "0.9615", "96.15% chance of ranking a fraud provider above a non-fraud provider — excellent discrimination"],
        ["Average Precision (AUPRC)", "0.7658", "Strong performance on the imbalanced fraud class specifically"],
        ["Fraud Recall", "81%", "The model correctly identifies 81% of all actual fraudulent providers"],
        ["Fraud Precision", "53%", "Of all flagged providers, 53% are confirmed fraud — acceptable for investigation workflows"],
        ["Overall Accuracy", "92%", "92% of all predictions are correct — strong general performance"],
        ["F1 Score (Fraud class)", "0.63", "Harmonic mean of precision and recall for the fraud class"],
    ],
    col_widths=[2.2, 1.2, 3.6]
)

add_body(doc, "Why Recall Is the Priority Metric", bold=True, size=12)
add_body(doc, (
    "In fraud detection, missing a fraudulent provider (False Negative) is far more costly than "
    "generating a false alarm on a legitimate provider (False Positive). A missed fraudster "
    "continues to collect fraudulent payments — potentially millions of dollars. A false alarm "
    "results only in an unnecessary audit, which is a manageable cost. Therefore, maximizing "
    "Recall (catching as many fraudsters as possible) is the primary objective."
))

add_body(doc, "Confusion Matrix Summary (Validation Set)", bold=True, size=12)
make_table(doc,
    ["", "Predicted: Not Fraud", "Predicted: Fraud"],
    [
        ["Actual: Not Fraud", "True Negatives (TN) — Correctly cleared", "False Positives (FP) — Unnecessary audits"],
        ["Actual: Fraud", "False Negatives (FN) — Missed fraudsters (costly!)", "True Positives (TP) — Fraudsters caught"],
    ],
    col_widths=[2.0, 2.5, 2.5]
)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════
# SECTION 9 — FINAL PREDICTIONS
# ══════════════════════════════════════════════════════════════════

add_heading(doc, "9. Final Predictions on Unseen Data", level=1)

add_body(doc, (
    "After training on the full training dataset (with SMOTE), the final XGBoost model was applied "
    "to the 1,353 providers in the unseen test set. The same data cleaning and feature engineering "
    "pipeline used for training data was applied identically to the test data."
))

add_body(doc, "Prediction Summary", bold=True, size=12)
make_table(doc,
    ["Metric", "Value"],
    [
        ["Total providers in unseen dataset", "1,353"],
        ["Predicted as Fraudulent (Yes)", "187 (13.8%)"],
        ["Predicted as Legitimate (No)", "1,166 (86.2%)"],
        ["Submission file name", "SupriyaSPoojary_Submission.csv"],
        ["Submission columns", "Provider | Probability | Predicted Class as per your best model"],
    ],
    col_widths=[3.5, 3.5]
)

add_body(doc, "Sample Predictions", bold=True, size=12)
make_table(doc,
    ["Provider", "Probability", "Predicted Class as per your best model"],
    [
        ["PRV51002", "0.0047", "No"],
        ["PRV55871", "0.8923", "Yes"],
        ["PRV52341", "0.1234", "No"],
        ["PRV53912", "0.7654", "Yes"],
        ["PRV54100", "0.0312", "No"],
    ],
    col_widths=[2.0, 2.0, 3.0]
)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════
# SECTION 10 — BUSINESS RECOMMENDATIONS
# ══════════════════════════════════════════════════════════════════

add_heading(doc, "10. Business Recommendations", level=1)

recommendations = [
    (
        "1. Implement a Risk-Tiered Investigation Framework",
        "Providers with fraud probability above 0.80 should be placed on an immediate hold — all "
        "pending claims paused and referred to the Special Investigation Unit (SIU). Providers "
        "scoring between 0.50 and 0.80 should be placed on an enhanced monitoring watchlist with "
        "additional manual review before payment. Providers below 0.50 proceed through normal "
        "payment processing."
    ),
    (
        "2. Integrate Real-Time Fraud Scoring into the Claims Pipeline",
        "The current model produces batch predictions. For maximum financial protection, the model "
        "should be deployed as a real-time API that scores each incoming claim before payment is "
        "released. This prevents fraudulent payments rather than recovering them after the fact."
    ),
    (
        "3. Establish a Monthly Model Retraining Cycle",
        "Fraudsters adapt their strategies over time. The model must be retrained monthly with newly "
        "labeled data to capture evolving patterns. A champion-challenger framework ensures the "
        "current production model is always compared against newly trained candidates."
    ),
    (
        "4. Investigate Physician Network Clusters",
        "The feature analysis reveals that fraudulent providers use unusually large networks of "
        "physician IDs. Investigators should map providers that share the same physician identifiers "
        "and audit them as potential fraud rings rather than isolated bad actors."
    ),
    (
        "5. Flag Providers with Disproportionate Chronic Condition Patient Bases",
        "Providers whose patient base includes a significantly higher-than-average proportion of "
        "patients with Alzheimer's disease or other cognitive impairments warrant enhanced scrutiny. "
        "These patients are unable to review and dispute their bills, making them prime targets for "
        "fraudulent billing."
    ),
    (
        "6. Deploy a Provider Risk Dashboard",
        "The Streamlit application built for this project provides a foundation for a production "
        "fraud monitoring dashboard. Insurance company investigators should have access to real-time "
        "provider risk scores, trend analysis, and one-click audit initiation."
    ),
    (
        "7. Apply Geographic Risk Weighting",
        "Certain states and counties historically show higher fraud rates. Incorporating regional "
        "risk indices into the model score can improve accuracy and help regional investigation teams "
        "prioritize workloads."
    ),
    (
        "8. Create a Confirmed Fraud Feedback Loop",
        "When investigators confirm or clear a flagged provider, those outcomes should be fed back "
        "as new training labels. This continuously improves model accuracy and ensures it learns "
        "from real investigation decisions."
    ),
    (
        "9. Conduct Annual ROI Analysis of the Fraud Detection System",
        "Track the total fraudulent claim amounts recovered or prevented versus the cost of "
        "running the ML system and investigation team. This justifies continued investment and "
        "identifies where additional resources would generate the most return."
    ),
    (
        "10. Combine Machine Learning with Rule-Based Fraud Indicators",
        "While the ML model is highly effective, certain deterministic rules (e.g., billing for "
        "a deceased patient, duplicate ClaimID submissions) should trigger immediate flags "
        "independent of the model score. A hybrid system combining ML scores and hard rules "
        "provides the strongest overall fraud detection."
    ),
]

for title, body in recommendations:
    add_body(doc, title, bold=True, size=11, color="1F3864")
    add_body(doc, body)
    doc.add_paragraph()

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════
# SECTION 11 — STREAMLIT APPLICATION
# ══════════════════════════════════════════════════════════════════

add_heading(doc, "11. Streamlit Application", level=1)

add_body(doc, (
    "An interactive web dashboard was built using Streamlit to make the model accessible to "
    "non-technical stakeholders including insurance investigators, fraud analysts, and management. "
    "The application is publicly accessible at the link below."
))

p_app2 = doc.add_paragraph()
p_app2.add_run("Live Application URL:  ").bold = True
add_hyperlink(p_app2, STREAMLIT_URL, STREAMLIT_URL)
doc.add_paragraph()

make_table(doc,
    ["Dashboard Page", "Features Available"],
    [
        ["Overview", "Project summary, business problem, methodology diagram, technology stack"],
        ["EDA & Insights", "Interactive EDA charts, correlation heatmap, fraud distribution plots"],
        ["Model Results", "Cross-validation comparison chart, confusion matrix, ROC curve, PR curve"],
        ["Live Prediction", "Manual input form with real-time fraud probability gauge and risk level badge"],
        ["Submission", "Fraud distribution chart, downloadable prediction CSV file"],
        ["Business Insights", "10 data-driven recommendations with impact ratings and explanations"],
        ["Documentation", "Technical deep-dive tabs covering cleaning, features, modeling, and run guide"],
    ],
    col_widths=[2.0, 5.0]
)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════
# SECTION 12 — PROJECT LINKS
# ══════════════════════════════════════════════════════════════════

add_heading(doc, "12. Project Links and Deliverables", level=1)

make_table(doc,
    ["Deliverable", "Description", "Link / File Name"],
    [
        ["GitHub Repository", "Full source code, notebooks, pipeline scripts, and model artifacts", GITHUB_URL],
        ["Live Streamlit App", "Interactive fraud detection dashboard accessible publicly", STREAMLIT_URL],
        ["Submission CSV", "Final predictions on 1,353 unseen providers with probability and class", "SupriyaSPoojary_Submission.csv"],
        ["HTML Notebook", "Fully executed Jupyter notebook exported to HTML with all results", "Healthcare_Provider_Fraud_Detection.html"],
        ["Custom Report", "Styled analysis report with all EDA and evaluation charts embedded", "FraudDetection_Notebook.html"],
        ["Word Document", "This document — full analysis, approach, and recommendations", "SupriyaSPoojary_Analysis_Report.docx"],
    ],
    col_widths=[1.8, 2.5, 2.7]
)

# Link paragraphs
p_g = doc.add_paragraph()
p_g.add_run("GitHub: ").bold = True
add_hyperlink(p_g, GITHUB_URL, GITHUB_URL)

p_s = doc.add_paragraph()
p_s.add_run("Streamlit App: ").bold = True
add_hyperlink(p_s, STREAMLIT_URL, STREAMLIT_URL)

doc.add_paragraph()

# ══════════════════════════════════════════════════════════════════
# SECTION 13 — CONCLUSION
# ══════════════════════════════════════════════════════════════════

add_heading(doc, "13. Conclusion", level=1)

add_body(doc, (
    "This project successfully developed an end-to-end Machine Learning pipeline for Healthcare "
    "Provider Fraud Detection. Starting from raw claim-level data across four datasets, the pipeline "
    "performs comprehensive data cleaning, engineers 56 provider-level behavioral features, selects "
    "54 of the most informative features, and trains an XGBoost classifier enhanced with SMOTE "
    "oversampling."
))
doc.add_paragraph()
add_body(doc, (
    "The final model achieves a ROC-AUC of 0.9615 on the validation set and correctly identifies "
    "81% of fraudulent providers — a performance level suitable for deployment in a production "
    "insurance fraud detection system. The model predicted 187 out of 1,353 unseen providers as "
    "potentially fraudulent, providing a concrete action list for investigation teams."
))
doc.add_paragraph()
add_body(doc, (
    "The accompanying Streamlit dashboard makes these predictions accessible to non-technical "
    "stakeholders, enabling real-time fraud probability scoring and supporting data-driven "
    "investigation decisions. The business recommendations outlined in this document provide a "
    "roadmap for integrating this system into production workflows and maximizing its financial "
    "impact for the insurance organization."
))

doc.add_paragraph()
doc.add_paragraph()
p_final = doc.add_paragraph()
p_final.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_final = p_final.add_run("Submitted by: Supriya S Poojary")
r_final.font.bold = True
r_final.font.size = Pt(13)
r_final.font.color.rgb = RGBColor(31, 56, 100)

# ── Save ─────────────────────────────────────────────────────────────────────
doc.save(OUTPUT_PATH)
print(f"Document saved successfully: {OUTPUT_PATH}")
